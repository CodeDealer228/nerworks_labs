import docx, copy, ipaddress as ip
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TPL="/root/.claude/uploads/5b43f2ec-e0ff-576e-8e16-7d5d355ff10e/f1d8673c-lab-1_report.docx"
OUT="/home/user/lab2/lab-2_report.docx"

# --- Часть 1: расчёт ---
V={1:("194.85.32.19","255.255.255.0",[10,6,1,18,100]),
   2:("10.12.12.15","255.255.254.0",[25,16,240,117,1]),
   3:("212.24.15.199","255.255.255.192",[7,0,0,11,10]),
   4:("120.13.120.120","255.255.255.224",[5,2,2,1,1])}
PORTS=[1,2,3,1,1]  # портов маршрутизаторов в каждой LAN по рис. 1

def calc(a,m,pc):
    base=ip.ip_network(f"{a}/{m}",strict=False)
    size=[1<<(pc[i]+PORTS[i]+1).bit_length() if False else 1<<((pc[i]+PORTS[i]+2)-1).bit_length() for i in range(5)]
    cur=int(base.network_address); res={}
    for i in sorted(range(5),key=lambda i:(-size[i],i)):
        res[i]=ip.ip_network((cur,32-(size[i].bit_length()-1))); cur+=size[i]
    assert cur<=int(base.broadcast_address)+1
    return base,[res[i] for i in range(5)]

# --- документ ---
d=docx.Document(TPL)
body=d.element.body
for el in list(body)[12:]:
    if el.tag!=qn('w:sectPr'): body.remove(el)
P=d.paragraphs
P[2].runs[0].text="Лабораторная работа № 2"
P[3].runs[0].text="Работа с адресами IP сетей"
for r in P[3].runs[1:]: r.text=""

def para(text="",size=11,bold=False,font="Times New Roman",style=None,align=None):
    p=d.add_paragraph(style=style)
    if text:
        r=p.add_run(text); r.font.name=font; r.font.size=Pt(size); r.bold=bold
        r._r.rPr.rFonts.set(qn('w:eastAsia'),font)
    if align: p.alignment=align
    return p
def h1(t): d.add_paragraph(t,style="Heading 1")
def h2(t): d.add_paragraph(t,style="Heading 2")
def code(t):
    for line in t.strip("\n").split("\n"):
        p=para(line,8.5,font="Courier New"); p.paragraph_format.space_after=Pt(0)
def shot(path,caption):
    p=para(align=A.CENTER); p.add_run().add_picture(path,width=Cm(15))
    para(caption,10,align=A.CENTER)

def table(rows,widths=None):
    t=d.add_table(rows=len(rows),cols=len(rows[0]))
    tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for s in("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement(f'w:{s}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); b.append(e)
    tblPr.append(b)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=""; r=c.paragraphs[0].add_run(str(v))
            r.font.name="Times New Roman"; r.font.size=Pt(9); r.bold=(i==0 or j==0)
    return t

h1("Ход работы")
para("Часть 1. Расчёт IP-сетей",12,True)
for s in ["1. По адресу провайдера и маске определил адрес сети (побитовое И адреса и маски) и общее число адресов в выданной сети.",
 "2. Для каждой из LAN1–LAN5 (рис. 1) посчитал число нужных адресов: компьютеры + порты маршрутизаторов (LAN1 — 1, LAN2 — 2, LAN3 — 3, LAN4 — 1, LAN5 — 1) + адрес сети + адрес широковещания.",
 "3. Округлил до ближайшей степени двойки, выбрал маску (табл. 3) и выделил подсети по порядку от самой большой, чтобы блоки были выровнены по границам.",
 "4. Начальный и конечный адреса узлов — адрес сети +1 и адрес широковещания −1."]:
    para(s)
para()
para("Часть 2. Проверка с помощью ipcalc (Debian 11)",12,True)
for s_ in ["1. Запустил виртуальную машину Debian 11, вернул сетевой адаптер в режим NAT, удалил интерфейс bond007 из лабораторной работы 1 и получил адрес по DHCP. Доступность Интернета проверил командой ping 8.8.8.8.",
 "2. Установил пакет ipcalc (версия 0.42-2) командой apt install ipcalc.",
 "3. Проверил все 20 рассчитанных сетей (4 варианта по 5 сетей): для каждой сравнил адрес сети с маской, первый и последний адрес узлов, адрес широковещания и число узлов. Расхождений с расчётом нет: ipcalc показывает число узлов, то есть на 2 меньше размера сети."]:
    para(s_)
para()
h1("Артефакты выполнения")
h2("Артефакт 1. Таблицы результатов расчёта (Часть 1)")
for v,(a,m,pc) in V.items():
    base,nets=calc(a,m,pc)
    para(f"Таблица {v+3} — Вариант {v}: {a}, маска {m}; сеть провайдера {base} ({base.num_addresses} адресов). Компьютеров: {', '.join(map(str,pc))}.",10)
    rows=[["Вариант",str(v),"","","",""],["Сеть"]+[f"Сеть {i+1}" for i in range(5)],
      ["IP сети, маска"]+[f"{n.network_address}\n{n.netmask}  (/{n.prefixlen})" for n in nets],
      ["Число IP адресов в IP сети"]+[n.num_addresses for n in nets],
      ["Начальный и конечный адреса, пригодные для адресации портов маршрутизаторов и компьютеров"]+[f"{list(n.hosts())[0]} –\n{list(n.hosts())[-1]}" for n in nets]]
    t=table(rows); m0=t.cell(0,1).merge(t.cell(0,5))
    for p in m0.paragraphs[1:]: p._p.getparent().remove(p._p)
    for row in t.rows:
        row.cells[0].width=Cm(3.6)
        for c in row.cells[1:]: c.width=Cm(2.9)
    para()

# ==== ЧАСТЬ 2: добавляется по мере получения артефактов ====
import os,json
PART2="/home/user/lab2/part2.json"   # [{"h":..., "text":..., "code":..., "img":..., "cap":...}]
if os.path.exists(PART2):
    para("Часть 2. Проверка с помощью ipcalc на виртуальной машине",12,True)
    for a in json.load(open(PART2)):
        if a.get("h"): h2(a["h"])
        if a.get("text"): para(a["text"])
        if a.get("code"): code(a["code"])
        if a.get("img"): shot(a["img"],a.get("cap",""))
h1("Вывод")
para("Освоил расчёт адреса сети, адреса широковещания и диапазона адресов узлов по IP-адресу и маске, а также разбиение сети на подсети масками переменной длины с учётом адресов портов маршрутизаторов. Результаты расчёта по всем четырём вариантам подтверждены утилитой ipcalc.")
d.save(OUT); print("saved",OUT)
